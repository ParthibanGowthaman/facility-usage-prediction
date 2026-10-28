"""Prediction review workbook: each test record with past bookings, the four predictions, the actuals,
match flags, plus a summary sheet."""
import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
GREEN, RED, HEAD = PatternFill("solid", fgColor="D9EAD3"), PatternFill("solid", fgColor="F4CCCC"), PatternFill("solid", fgColor="1F3864")


def _hm(minutes):
    m = int(round(minutes)) % 1440
    return f"{m // 60:02d}:{m % 60:02d}"


def _past(bookings, fac_names, n=3):
    """Text for each booking's previous n bookings, newest last: 'Gym / Mon / 07:00 / booked Sun 18:20'."""
    b = bookings.sort_values(["booking_ts", "booking_id"])
    line = (b.facility_id.map(fac_names) + " / " + b.usage_ts.dt.dayofweek.map(dict(enumerate(DAYS))) + " / "
            + b.usage_ts.dt.strftime("%H:%M") + " / booked " + b.booking_ts.dt.dayofweek.map(dict(enumerate(DAYS)))
            + " " + b.booking_ts.dt.strftime("%H:%M"))
    g = pd.DataFrame({"resident_id": b.resident_id, "line": line})
    cols = [g.groupby("resident_id").line.shift(i) for i in range(n, 0, -1)]
    joined = pd.concat(cols, axis=1).apply(lambda r: "\n".join(x for x in r if isinstance(x, str)), axis=1)
    return joined.where(joined != "", "(no history)").set_axis(b.booking_id)


def review_table(X, P, h, bookings, fac_names):
    past = _past(bookings, fac_names)
    t = pd.DataFrame({"Booking ID": X.booking_id.values, "Resident": X.resident_id.values,
                      "Past bookings (last 3)": past.reindex(X.booking_id).values})
    fn = lambda s: s.map(lambda i: fac_names[f"F{int(i) + 1}"])
    t["Predicted facility"] = fn(P.p_fac).values
    t["Predicted day"] = P.p_dow.map(dict(enumerate(DAYS))).values
    t["Predicted hour"] = P.p_hour.map(lambda v: f"{int(v):02d}:00").values
    t["Predicted nudge"] = (P.p_nudge_dow.map(dict(enumerate(DAYS))) + " " + P.p_clock.map(_hm)).values
    t["Actual facility"] = fn(X.y_fac).values
    t["Actual day"] = X.y_dow.map(dict(enumerate(DAYS))).values
    t["Actual hour"] = X.y_hour.map(lambda v: f"{int(v):02d}:00").values
    t["Actual booked at"] = (X.y_nudge_dow.map(dict(enumerate(DAYS))) + " " + X.y_clock.map(_hm)).values
    for c, k in (("Facility", "facility"), ("Day", "day"), ("Hour", "hour"), ("Nudge", "nudge")):
        t[f"{c} match"] = np.where(h[k].values, "YES", "NO")
    t["Score"] = h.k_of_4.map(lambda k: f"{k} of 4").values
    t["All 4 match"] = np.where(h.all_4.values, "YES", "NO")
    return t


def write_workbook(path, table, summary_blocks):
    """summary_blocks: list of (title, DataFrame). Written to a 'Summary' sheet, then 'Predictions'."""
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        row = 0
        for title, df in summary_blocks:
            pd.DataFrame([[title]]).to_excel(xw, "Summary", startrow=row, index=False, header=False)
            df.to_excel(xw, "Summary", startrow=row + 1)
            row += len(df) + 4
        table.to_excel(xw, "Predictions", index=False)
        ws, summ = xw.sheets["Predictions"], xw.sheets["Summary"]
        for cell in ws[1]:
            cell.fill, cell.font = HEAD, Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for i, col in enumerate(table.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = 44 if col.startswith("Past") else 17
        flag_cols = [i for i, c in enumerate(table.columns, 1) if c.endswith("match")]
        for r in ws.iter_rows(min_row=2):
            r[2].alignment = Alignment(wrap_text=True, vertical="top")
            for i in flag_cols:
                r[i - 1].fill = GREEN if r[i - 1].value == "YES" else RED
        ws.freeze_panes, ws.auto_filter.ref = "C2", ws.dimensions
        summ.column_dimensions["A"].width = 30
        for i in range(2, 12):
            summ.column_dimensions[get_column_letter(i)].width = 16
        for r in summ.iter_rows():
            if r[0].value and all(c.value is None for c in r[1:3]):
                r[0].font = Font(bold=True, size=12)
